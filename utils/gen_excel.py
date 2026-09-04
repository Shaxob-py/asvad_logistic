import datetime

import msoffcrypto

from db.adler import TypeCompany, GroupEntry


def get_inkassatsiya_row(file_bytes: bytes, password: str, target_date: datetime.date,
                         group_entry,
                         date_col=1, statya_col=2, amount_col=6, balance_col=7):
    decrypted = io.BytesIO()
    office_file = msoffcrypto.OfficeFile(io.BytesIO(file_bytes))
    office_file.load_key(password=password)
    office_file.decrypt(decrypted)
    decrypted.seek(0)

    if group_entry.company == TypeCompany.GATTER:
        target_date -= datetime.timedelta(days=1)

    wb = openpyxl.load_workbook(decrypted, data_only=True, read_only=True)
    ws = wb.worksheets[0]

    prev_balance = None
    rows_by_date: dict[datetime.date, list[dict]] = {}

    for row in ws.iter_rows(values_only=True):
        row_date = row[date_col]
        if row_date is None:
            continue
        cell_date = row_date.date() if isinstance(row_date, datetime.datetime) else row_date

        amount = row[amount_col]
        balance = row[balance_col]
        statya = (row[statya_col] or "").strip().lower() if statya_col is not None else None

        rows_by_date.setdefault(cell_date, []).append({
            "amount": amount,
            "balance_before": prev_balance,
            "balance_after": balance,
            "statya": statya,
        })

        if balance is not None:
            prev_balance = balance

    def find_matching_row(date, skip_income=False):
        for r in rows_by_date.get(date, []):
            if skip_income and r["statya"] == "приход":
                continue
            if r["amount"] is not None and r["amount"] == group_entry.inkassatsiya_amount:
                return r
        return None

    found = find_matching_row(target_date)
    if found is not None:
        return found["balance_before"], group_entry.inkassatsiya_amount, found["balance_after"]

    # 2) ищем на target_date - 1, пропуская строки "Приход"
    prev_date = target_date - datetime.timedelta(days=1)
    found = find_matching_row(prev_date, skip_income=True)
    if found is not None:
        return found["balance_before"], group_entry.inkassatsiya_amount, found["balance_after"]

    prev_rows = rows_by_date.get(prev_date, [])
    if prev_rows:
        balance_before = prev_rows[-1]["balance_after"]  # последний известный остаток
        balance_after = balance_before - group_entry.inkassatsiya_amount
        return balance_before, group_entry.inkassatsiya_amount, balance_after

    return None, group_entry.inkassatsiya_amount, None


# generate Excel

from datetime import date

import openpyxl
from openpyxl.styles import Border, Side

CANONICAL_CITIES = [
    "Самарканд", "Карши", "Наманган", "Андижон", "Термез", "Алмалык",
    "Фергана", "Гулистон", "Бухара", "Ургенч", "Нукус", "Коканд",
    "Навои", "Джиззах",
]


def generate_summary_excel(target_date: date, entries: list) -> bytes:
    data = {"ADLER": {}, "GATTER": {}}
    for e in entries:
        company = e.company.value if hasattr(e.company, "value") else e.company
        data[company][e.region] = e

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Сводка"

    header_fill = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
    row_fill = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")
    bold = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center")
    right = Alignment(horizontal="right", vertical="center")
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    red_font = Font(color="FF0000", bold=True)

    date_str = target_date.strftime("%d.%m.%Y")

    def write_block(cols, company_name, headers):
        first, last = cols[0], cols[-1]
        ws.merge_cells(f"{first}1:{last}1")
        ws[f"{first}1"] = date_str
        ws[f"{first}1"].font = bold
        ws[f"{first}1"].alignment = center
        ws[f"{first}1"].fill = header_fill

        ws.merge_cells(f"{first}2:{last}2")
        ws[f"{first}2"] = company_name
        ws[f"{first}2"].font = bold
        ws[f"{first}2"].alignment = center
        ws[f"{first}2"].fill = header_fill

        for c, text in zip(cols, headers):
            cell = ws[f"{c}3"]
            cell.value = text
            cell.font = bold
            cell.alignment = center
            cell.fill = header_fill
            cell.border = border

    adler_cols = "BCDEF"
    write_block(adler_cols, "ADLER",
                ["№", "Регион", "AGD остатка", "AGD инкассация", "Остаток сумма после инкассация AGD"])

    gatter_cols = "IJKLM"
    write_block(gatter_cols, "GATTER",
                ["№", "Регион", "GG остатка", "GG инкассация", "Остаток сумма после инкассация GG"])

    start_row = 4
    for i, region in enumerate(CANONICAL_CITIES):
        row = start_row + i
        for company, cols in (("ADLER", adler_cols), ("GATTER", gatter_cols)):
            entry = data[company].get(region)
            ws[f"{cols[0]}{row}"] = i + 1
            ws[f"{cols[1]}{row}"] = region
            if entry:
                ws[f"{cols[2]}{row}"] = entry.balance_before
                cell_ink = ws[f"{cols[3]}{row}"]
                cell_ink.value = entry.inkassatsiya_amount
                cell_ink.font = red_font
                ws[f"{cols[4]}{row}"] = entry.balance_after

            for col in cols:
                cell = ws[f"{col}{row}"]
                cell.border = border
                cell.fill = row_fill
                if col in (cols[0], cols[1]):
                    cell.alignment = center
                else:
                    cell.alignment = right
                    if cell.value is not None:
                        cell.number_format = "#,##0"

    widths = {"A": 3, "B": 4, "C": 14, "D": 15, "E": 15, "F": 32,
              "G": 3, "H": 3,
              "I": 4, "J": 14, "K": 15, "L": 15, "M": 32}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


# Стили под каждую компанию: цвет заголовка + цвет шрифта заголовка
COMPANY_STYLES = {
    "ADLER": {
        "header_fill": "FF1F4E78",   # тёмно-синий
        "header_font": "FFFFFFFF",   # белый текст
    },
    "GATTER": {
        "header_fill": "FF7F6000",   # тёмно-жёлтый/охра
        "header_font": "FFFFFFFF",
    },
}

DEFAULT_STYLE = {
    "header_fill": "FF808080",  # серый — на случай новой компании
    "header_font": "FFFFFFFF",
}
import io
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# Цвет строки для каждой компании (заливка данных, не только заголовка)
COMPANY_ROW_COLORS = {
    "ADLER": "FFD9E1F2",   # светло-синий
    "GATTER": "FFFCE4B4",  # светло-жёлтый/охра
}

DEFAULT_ROW_COLOR = "FFEFEFEF"  # серый — на случай новой/незнакомой компании

HEADER_FILL = "FF404040"
HEADER_FONT = "FFFFFFFF"


def generate_inkassatsiya_excel(data: list[GroupEntry]) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Инкассация"

    # Заголовки
    ws.append([
        "Компания",
        "Дата",
        "Регион",
        "Сумма инкассации",
    ])

    header_fill = PatternFill(start_color=HEADER_FILL, end_color=HEADER_FILL, fill_type="solid")
    header_font = Font(bold=True, color=HEADER_FONT)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Данные — все компании на одном листе, но с разной заливкой строки
    for item in data:
        company = item.company.value  # ADLER / GATTER
        row_color = COMPANY_ROW_COLORS.get(company, DEFAULT_ROW_COLOR)
        fill = PatternFill(start_color=row_color, end_color=row_color, fill_type="solid")

        ws.append([
            company,
            item.date,
            item.region,
            item.inkassatsiya_amount,
        ])

        for cell in ws[ws.max_row]:
            cell.fill = fill

    # Автоматическая ширина колонок
    for column in ws.columns:
        max_length = 0
        for cell in column:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))
        column_letter = get_column_letter(column[0].column)
        ws.column_dimensions[column_letter].width = max_length + 2

    # Закрепить заголовок
    ws.freeze_panes = "A2"

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return output